"""Generate the small, synthetic document-parser fixtures committed here.

Run this script in the constrained parser virtual environment.  It avoids
system clocks in document metadata and rewrites Office ZIP files with fixed
timestamps so consecutive runs have identical SHA-256 hashes.
"""

from __future__ import annotations

import io
import os
import re
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import fitz
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches as PptInches


ROOT = Path(__file__).resolve().parent
FIXED_TIME = datetime(2026, 1, 1, tzinfo=timezone.utc)
ZIP_TIME = (2026, 1, 1, 0, 0, 0)
PDF_TIME = "D:20260101000000+00'00'"
OOXML_MODIFIED = (
    b'<dcterms:modified xmlns:dcterms="http://purl.org/dc/terms/" '
    b'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" '
    b'xsi:type="dcterms:W3CDTF">2026-01-01T00:00:00Z</dcterms:modified>'
)


def chinese_font_path() -> Path:
    """Find a common CJK font; fail rather than silently writing tofu glyphs."""
    candidates = (
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "NotoSansSC-VF.ttf",
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "msyh.ttc",
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"),
        Path("/System/Library/Fonts/PingFang.ttc"),
    )
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    raise RuntimeError("A CJK font is required to generate OCR fixtures; install Noto Sans CJK.")


def chinese_font(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(chinese_font_path()), size=size)


def make_text_image(path: Path, first_line: str = "图片 OCR", second_line: str = "161803") -> None:
    image = Image.new("RGB", (1500, 500), "white")
    draw = ImageDraw.Draw(image)
    font = chinese_font(76)
    draw.text((70, 75), first_line, fill="black", font=font)
    draw.text((70, 245), second_line, fill="black", font=font)
    image.save(path, format="PNG", optimize=False, compress_level=9)


def make_pdf(path: Path, scanned: bool) -> None:
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    if scanned:
        canvas = Image.new("RGB", (1500, 700), "white")
        draw = ImageDraw.Draw(canvas)
        font = chinese_font(76)
        draw.text((70, 110), "Qavor 扫描 PDF", fill="black", font=font)
        draw.text((70, 340), "OCR 314159", fill="black", font=font)
        buffer = io.BytesIO()
        canvas.save(buffer, format="PNG", optimize=False, compress_level=9)
        page.insert_image(page.rect, stream=buffer.getvalue())
    else:
        page.insert_font(fontname="qavorcjk", fontfile=str(chinese_font_path()))
        page.insert_text((72, 90), "Qavor 数字 PDF", fontsize=20, fontname="qavorcjk")
        page.insert_text((72, 130), "结构化表格", fontsize=16, fontname="qavorcjk")
        page.draw_rect(fitz.Rect(72, 160, 440, 250), color=(0, 0, 0), width=0.8)
        page.draw_line((72, 190), (440, 190), color=(0, 0, 0), width=0.8)
        page.draw_line((230, 160), (230, 250), color=(0, 0, 0), width=0.8)
        page.insert_text((90, 182), "Year", fontsize=12, fontname="helv")
        page.insert_text((250, 182), "Value", fontsize=12, fontname="helv")
        page.insert_text((90, 222), "2026", fontsize=12, fontname="helv")
        page.insert_text((250, 222), "42", fontsize=12, fontname="helv")
    pdf.set_metadata(
        {
            "title": "Qavor synthetic parser fixture",
            "author": "Qavor",
            "creator": "Qavor fixture generator",
            "producer": "Qavor fixture generator",
            "creationDate": PDF_TIME,
            "modDate": PDF_TIME,
        }
    )
    pdf.subset_fonts()
    pdf.save(path, garbage=4, deflate=True, no_new_id=True)
    pdf.close()


def normalize_office_zip(path: Path) -> None:
    """Eliminate ZIP timestamps and ordering differences emitted by Office writers."""
    with zipfile.ZipFile(path, "r") as source:
        entries = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".zip", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with zipfile.ZipFile(temporary_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
            for filename, payload in sorted(entries):
                if filename == "docProps/core.xml":
                    payload = re.sub(
                        rb"<dcterms:modified[^>]*>[^<]*</dcterms:modified>",
                        OOXML_MODIFIED,
                        payload,
                    )
                info = zipfile.ZipInfo(filename, date_time=ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = 3
                info.external_attr = 0o600 << 16
                target.writestr(info, payload)
        temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)


def set_core_properties(properties: object) -> None:
    properties.author = "Qavor"  # type: ignore[attr-defined]
    properties.title = "Qavor synthetic parser fixture"  # type: ignore[attr-defined]
    properties.subject = "Synthetic CC0 fixture"  # type: ignore[attr-defined]
    properties.created = FIXED_TIME  # type: ignore[attr-defined]
    properties.modified = FIXED_TIME  # type: ignore[attr-defined]


def make_docx(path: Path) -> None:
    document = Document()
    set_core_properties(document.core_properties)
    document.add_heading("DOCX 正文", level=1)
    document.add_paragraph("这是可重复生成的 Office 解析样本。")
    image_path = ROOT / ".docx-image.png"
    try:
        make_text_image(image_path, "图片文字", "271828")
        document.add_picture(str(image_path), width=Inches(5.5))
        document.save(path)
    finally:
        image_path.unlink(missing_ok=True)
    normalize_office_zip(path)


def make_pptx(path: Path) -> None:
    presentation = Presentation()
    set_core_properties(presentation.core_properties)
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "PPTX 标题"
    slide.placeholders[1].text = "演示正文"
    slide.shapes.add_picture(str(ROOT / "text-image.png"), PptInches(1), PptInches(3), width=PptInches(4))
    presentation.save(path)
    normalize_office_zip(path)


def make_xlsx(path: Path) -> None:
    workbook = Workbook()
    workbook.properties.creator = "Qavor"
    workbook.properties.title = "Qavor synthetic parser fixture"
    workbook.properties.created = FIXED_TIME
    workbook.properties.modified = FIXED_TIME
    sheet = workbook.active
    sheet.title = "产品表"
    sheet.append(["产品", "数量"])
    sheet.append(["Qavor", 42])
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 12
    workbook.save(path)
    normalize_office_zip(path)


def main() -> None:
    (ROOT / "plain.txt").write_text("页眉：Qavor 文档\r\n\x01\r\n\r\n正文 42   \r\n\r\n页脚：第 1 页\r\n", encoding="utf-8")
    (ROOT / "structured.md").write_text(
        "# 结构化样本\n\n|产品|数量|\n|-|-|\n|Qavor|42|\n\n```go\nfunc main() {}\n```\n\n![图片 OCR](https://example.invalid/image.png)\n\n页脚：第 1 页\n",
        encoding="utf-8",
    )
    make_text_image(ROOT / "text-image.png")
    make_pdf(ROOT / "digital.pdf", scanned=False)
    make_pdf(ROOT / "scanned.pdf", scanned=True)
    make_docx(ROOT / "with-image.docx")
    make_pptx(ROOT / "slides.pptx")
    make_xlsx(ROOT / "table.xlsx")
    (ROOT / "corrupted.pdf").write_bytes(b"%PDF-1.7\nQavor intentionally corrupted test fixture\n%%EOF\x00trailing-garbage")


if __name__ == "__main__":
    main()
